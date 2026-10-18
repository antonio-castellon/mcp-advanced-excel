"""Regenerate the synthetic spreadsheets under tests/fixtures/.

Run from the repo root:

    python tests/fixtures/build_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

import xlwt
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent


def build() -> None:
    _csv_decimal()
    _xlsx_instrument()
    _xlsx_merged_header()
    _xlsx_formula()
    _xls_simple()


def _csv_decimal() -> None:
    # Semicolon CSV with decimal commas, a metadata block, and a footer note.
    text = (
        "Instrument: ICP-MS\n"
        "\n"
        "Sample;Result (mg/L);Flag\n"
        "A-1;0,125;ok\n"
        "A-2;1.250,5;ok\n"
        "Notes: end of report\n"
    )
    (HERE / "decimal_comma.csv").write_text(text, encoding="utf-8")


def _xlsx_instrument() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Results"
    sheet["A1"] = "Instrument: Agilent 1260"
    sheet["A2"] = "Method: HPLC-UV"
    sheet["A4"] = "Sample ID"
    sheet["B4"] = "Compound"
    sheet["C4"] = "Conc. (ug/mL)"
    sheet["D4"] = "Temp (K)"
    sheet["A5"] = "S-001"
    sheet["B5"] = "Caffeine"
    sheet["C5"] = "12,5"
    sheet["D5"] = "298,15"
    sheet["A6"] = "S-002"
    sheet["B6"] = "Caffeine"
    sheet["C6"] = "1.234,5"
    sheet["D6"] = 300
    sheet["A8"] = "End of report. Do not distribute."
    book.save(HERE / "instrument.xlsx")


def _xlsx_merged_header() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Wide"
    sheet["A1"] = "Project: Demo"
    sheet["A3"] = "Sample"
    sheet["B3"] = "Concentration"
    sheet.merge_cells("B3:C3")
    sheet["D3"] = "Temperature"
    sheet["A4"] = "ID"
    sheet["B4"] = "mg/L"
    sheet["C4"] = "ug/mL"
    sheet["D4"] = "degC"
    sheet["A5"] = "S-1"
    sheet["B5"] = 10
    sheet["C5"] = 10
    sheet["D5"] = 25
    book.save(HERE / "merged_header.xlsx")


def _xlsx_formula() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Calc"
    sheet["A1"] = "Sample"
    sheet["B1"] = "Reading"
    sheet["C1"] = "Sum"
    sheet["A2"] = "S-1"
    sheet["B2"] = 1
    sheet["C2"] = "=1+1"
    book.save(HERE / "formula.xlsx")


def _xls_simple() -> None:
    book = xlwt.Workbook()
    sheet = book.add_sheet("Results")
    sheet.write(0, 0, "Name")
    sheet.write(0, 1, "Mass (mg)")
    sheet.write(1, 0, "alpha")
    sheet.write(1, 1, "1,5")
    sheet.write(2, 0, "beta")
    sheet.write(2, 1, "2,0")
    book.save(str(HERE / "simple.xls"))


if __name__ == "__main__":
    build()
